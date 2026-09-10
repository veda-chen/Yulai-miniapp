from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
reference = Path(r'C:\Users\chenz\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx')
source = root / '羽来技术架构设计和开发排期_V0.1.md'
cover_image = root / '小程序头像.jpeg'
output = root / '羽来技术架构设计和开发排期_V0.1.docx'

doc = Document(reference)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

for style_name in ['normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['normal'].font.size = Pt(10.5)
doc.styles['normal'].paragraph_format.line_spacing = 1.2
doc.styles['normal'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].paragraph_format.keep_with_next = True

doc.core_properties.title = '羽来技术架构设计和开发排期 V0.1'
doc.core_properties.author = ''

cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_after = Pt(22)
picture = cover.add_run().add_picture(str(cover_image), width=Inches(5.55))
picture._inline.docPr.set('descr', '羽来羽毛球小程序形象图')
picture._inline.docPr.set('title', '羽来')
title = doc.add_paragraph('羽来技术架构设计和开发排期', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph('V1工程基线与十二周试点计划 V0.1')
subtitle.paragraph_format.space_after = Pt(4)
subtitle.runs[0].font.size = Pt(13)
meta = doc.add_paragraph('广东工业大学大学城校区体育馆试点  |  2026年9月8日')
meta.runs[0].font.size = Pt(9.5)
meta.runs[0].font.color.rgb = RGBColor(90, 90, 90)
doc.add_section(WD_SECTION.NEW_PAGE)

lines = source.read_text(encoding='utf-8').splitlines()
i = 2

def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)

def borders(table):
    tbl_pr = table._tbl.tblPr
    elem = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), 'D9D9D9')
        elem.append(tag)
    tbl_pr.append(elem)

def set_table_geometry(table, ratios, total_twips=9288):
    widths = [round(total_twips * ratio) for ratio in ratios[:-1]]
    widths.append(total_twips - sum(widths))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(total_twips))
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:type'), 'dxa')
    tbl_ind.set(qn('w:w'), '100')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in('w:tcW')
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                cell._tc.get_or_add_tcPr().insert(0, tc_w)
            tc_w.set(qn('w:type'), 'dxa')
            tc_w.set(qn('w:w'), str(width))

def add_text_with_links(paragraph, text_value):
    url_pattern = re.compile(r'https?://[^\s]+')
    pos = 0
    for match in url_pattern.finditer(text_value):
        if match.start() > pos:
            paragraph.add_run(text_value[pos:match.start()])
        url = match.group(0)
        part = doc.part
        r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)
        run = OxmlElement('w:r')
        r_pr = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '1F6B5C')
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        r_pr.append(color)
        r_pr.append(underline)
        run.append(r_pr)
        text_node = OxmlElement('w:t')
        text_node.text = '打开官方页面' if text_value.strip() == url else url
        run.append(text_node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        pos = match.end()
    if pos < len(text_value):
        paragraph.add_run(text_value[pos:])

while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            values = [x.strip() for x in lines[i].strip('|').split('|')]
            if not all(set(x) <= set('-: ') for x in values):
                rows.append(values)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.autofit = False
        borders(table)
        total = 6.45
        ratios = {2: [0.28, 0.72], 3: [0.23, 0.39, 0.38], 4: [0.16, 0.30, 0.36, 0.18]}.get(len(rows[0]), [1 / len(rows[0])] * len(rows[0]))
        for row_index, values in enumerate(rows):
            cells = table.add_row().cells
            for col_index, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Inches(total * ratios[col_index])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                margins(cell)
                cell.text = ''
                add_text_with_links(cell.paragraphs[0], value)
                if row_index == 0:
                    shade(cell, '234E43')
                elif row_index % 2 == 0:
                    shade(cell, 'F2F7F5')
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.name = 'Microsoft YaHei'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                        run.font.size = Pt(8.6)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            tr_pr = table.rows[-1]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement('w:cantSplit'))
            if row_index == 0:
                tr_pr.append(OxmlElement('w:tblHeader'))
        set_table_geometry(table, ratios)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if line.startswith('### '):
        heading = re.sub(r'^(\d+)\.(\d+)\s*', r'\1 \2 ', line[4:])
        doc.add_paragraph(heading, 'Heading 2')
    elif line.startswith('## '):
        heading = re.sub(r'^(\d+)\s*', r'\1 ', line[3:])
        doc.add_paragraph(heading, 'Heading 1')
    elif re.match(r'^\d+\.\s', line):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        number, value = re.match(r'^(\d+)\.\s*(.*)$', line).groups()
        paragraph.add_run(f'{number}. ')
        add_text_with_links(paragraph, value)
    elif not line.startswith('# '):
        paragraph = doc.add_paragraph()
        add_text_with_links(paragraph, line)
    i += 1

for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('羽来  技术架构设计和开发排期  V0.1')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(110, 110, 110)

doc.save(output)

check = Document(output)
all_text = '\n'.join(p.text for p in check.paragraphs) + '\n' + '\n'.join(
    ' | '.join(c.text for c in row.cells) for table in check.tables for row in table.rows
)
required = [
    '模块化单体', 'PostgreSQL', '广东工业大学大学城校区体育馆',
    '学校始终可选且可清空', '100并发请求不超容量', '2026年12月6日',
    '正式参与者可录入比分', '订阅提醒只能在用户主动选择后使用'
]
missing = [item for item in required if item not in all_text]
assert not missing, missing
assert len(check.sections) == 2
assert len(check.tables) >= 16
assert output.stat().st_size > 100000
print(f'Validated: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {len(check.sections)} sections, {output.stat().st_size} bytes')
