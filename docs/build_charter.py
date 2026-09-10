from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
src = root / '项目章程与V1需求范围_V0.1.md'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2.1)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(10.5 if name == 'Normal' else {'Title': 24, 'Heading 1': 16, 'Heading 2': 12}[name])
    style.font.color.rgb = RGBColor.from_string('000000')
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.2
doc.core_properties.title = '羽来项目章程和V1需求范围说明 V0.2'
doc.core_properties.author = ''
lines = src.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [s.strip() for s in lines[i].strip('|').split('|')]
            if not all(set(c) <= set('-: ') for c in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = 'Table Grid'
        for idx, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.bold = idx == 0
                if idx == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8F0F4')
                    cell._tc.get_or_add_tcPr().append(shade)
            trpr = table.rows[-1]._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if idx == 0:
                trpr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:].replace('·', ' ').replace('+', '和'), 'Title')
    elif line.startswith('### '):
        heading = re.sub(r'^(\d+)\.(\d+)\s*', r'\1 \2 ', line[4:])
        doc.add_paragraph(heading.replace('+', '和'), 'Heading 2')
    elif line.startswith('## '):
        heading = re.sub(r'^(\d+)\.\s*', r'\1 ', line[3:])
        doc.add_paragraph(heading.replace('+', '和'), 'Heading 1')
    else:
        doc.add_paragraph(line)
    i += 1
out = root / '羽来项目章程与V1需求范围说明_V0.2.docx'
doc.save(out)
check = Document(out)
assert len(check.tables) == 6, len(check.tables)
assert all(len(t.rows) > 1 for t in check.tables)
assert '11 待决策项与下一步' in [p.text for p in check.paragraphs]
print(f'Validated: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {out.stat().st_size} bytes')
