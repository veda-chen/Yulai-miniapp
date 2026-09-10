from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
reference = Path(r'C:\Users\chenz\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx')
source = root / '羽来PRD_V0.1.md'
cover_image = root / '小程序头像.jpeg'
output = root / '羽来产品需求文档_V0.1.docx'

doc = Document(reference)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

for style_name in ['normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['normal'].font.size = Pt(10.5)
doc.styles['normal'].paragraph_format.line_spacing = 1.2
doc.styles['normal'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].paragraph_format.keep_with_next = True

doc.core_properties.title = '羽来产品需求文档 组局主流程和可选现场功能 V0.1'
doc.core_properties.author = ''

cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_after = Pt(22)
picture = cover.add_run().add_picture(str(cover_image), width=Inches(5.55))
doc_pr = picture._inline.docPr
doc_pr.set('descr', '羽来羽毛球小程序形象图')
doc_pr.set('title', '羽来')
title = doc.add_paragraph('羽来产品需求文档', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph('组局主流程和可选现场功能 V0.1')
subtitle.paragraph_format.space_after = Pt(4)
subtitle.runs[0].font.size = Pt(13)
meta = doc.add_paragraph('长期运营需求草案  |  2026年9月8日')
meta.runs[0].font.size = Pt(9.5)
meta.runs[0].font.color.rgb = RGBColor(90, 90, 90)
doc.add_section(WD_SECTION.NEW_PAGE)

lines = source.read_text(encoding='utf-8').splitlines()
i = 4

def set_cell_shading(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)

def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), 'D9D9D9')
        borders.append(tag)
    tbl_pr.append(borders)

while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            values = [x.strip() for x in lines[i].strip('|').split('|')]
            if not all(set(x) <= set('-: ') for x in values):
                rows.append(values)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.autofit = False
        set_table_borders(table)
        total = 6.45
        ratios = {2: [0.31, 0.69], 3: [0.20, 0.33, 0.47], 4: [0.15, 0.28, 0.39, 0.18]}.get(len(rows[0]), [1 / len(rows[0])] * len(rows[0]))
        for row_index, values in enumerate(rows):
            cells = table.add_row().cells
            for col_index, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Inches(total * ratios[col_index])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                cell.text = value
                if row_index == 0:
                    set_cell_shading(cell, '234E43')
                elif row_index % 2 == 0:
                    set_cell_shading(cell, 'F2F7F5')
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.08
                    for run in paragraph.runs:
                        run.font.name = 'Microsoft YaHei'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            tr_pr = table.rows[-1]._tr.get_or_add_trPr()
            cant_split = OxmlElement('w:cantSplit')
            tr_pr.append(cant_split)
            if row_index == 0:
                tr_pr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if line.startswith('### '):
        heading = re.sub(r'^(\d+)\.(\d+)\s*', r'\1 \2 ', line[4:])
        doc.add_paragraph(heading, 'Heading 2')
    elif line.startswith('## '):
        heading = re.sub(r'^(\d+)\s*', r'\1 ', line[3:])
        doc.add_paragraph(heading, 'Heading 1')
    elif not line.startswith('# '):
        doc.add_paragraph(line)
    i += 1

doc.save(output)
check = Document(output)
text = '\n'.join(p.text for p in check.paragraphs) + '\n' + '\n'.join(
    ' | '.join(c.text for c in row.cells) for table in check.tables for row in table.rows
)
required = ['V1首先解决', '广东工业大学大学城校区体育馆', '学校 | 否', '正式参与者都可以录入', 'U12']
assert all(item in text for item in required), {item: item in text for item in required}
assert len(check.sections) == 2
assert len(check.tables) >= 10
print(f'Validated: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {len(check.sections)} sections, {output.stat().st_size} bytes')
